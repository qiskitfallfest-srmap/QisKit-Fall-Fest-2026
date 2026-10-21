"""
Sync Competition Evaluations from XLSX back to Supabase
-------------------------------------------------------
Reads QFF2026_Online_Competitions_Evaluation_Master.xlsx and upserts
evaluations into public.submission_evaluations.

Usage:
    python tools/sync_competition_evaluations_from_xlsx.py [--file path/to/file.xlsx] [--dry-run]
"""

import os
import sys
import argparse
from datetime import datetime, timezone
import openpyxl
import requests
from dotenv import load_dotenv

load_dotenv('.env.local')

SUPABASE_URL = os.environ.get('NEXT_PUBLIC_SUPABASE_URL')
SERVICE_KEY = os.environ.get('SUPABASE_SERVICE_ROLE_KEY')

HEADERS = {
    'apikey': SERVICE_KEY,
    'Authorization': f'Bearer {SERVICE_KEY}',
    'Content-Type': 'application/json',
    'Prefer': 'resolution=merge-duplicates'
}

VERDICT_MAP = {
    'amazing': ('finalist', True),
    'shortlist': ('shortlisted', True),
    'reject': ('rejected', False),
    'under review': ('under_review', False),
}

def sync_from_xlsx(file_path='QFF2026_Online_Competitions_Evaluation_Master.xlsx', dry_run=False):
    if not os.path.exists(file_path):
        print(f"Error: File '{file_path}' not found.")
        sys.exit(1)

    print(f"Loading workbook from {file_path}...")
    wb = openpyxl.load_workbook(file_path, data_only=True)

    sheets_to_process = [
        ('Tech Reels', 'reels'),
        ('Digital Poster', 'poster'),
        ('Technical Essay', 'essay'),
    ]

    records_to_upsert = []

    for sheet_name, category in sheets_to_process:
        if sheet_name not in wb.sheetnames:
            print(f"Warning: Sheet '{sheet_name}' not found in workbook, skipping.")
            continue

        ws = wb[sheet_name]
        start_row = 5
        max_row = ws.max_row

        for row in range(start_row, max_row + 1):
            sub_id = ws.cell(row=row, column=16).value
            if not sub_id or not isinstance(sub_id, str) or len(sub_id) < 10:
                continue

            verdict_raw = str(ws.cell(row=row, column=2).value or '').strip()
            score_raw = ws.cell(row=row, column=3).value
            evaluator_name = str(ws.cell(row=row, column=13).value or '').strip()
            comments_raw = str(ws.cell(row=row, column=14).value or '').strip()

            v_lower = verdict_raw.lower()
            if v_lower not in VERDICT_MAP:
                status, is_next_round = 'under_review', False
            else:
                status, is_next_round = VERDICT_MAP[v_lower]

            score_val = None
            if score_raw is not None and str(score_raw).strip() != '':
                try:
                    score_val = float(score_raw)
                except ValueError:
                    pass

            comments_list = []
            if comments_raw:
                comments_list.append({
                    'id': f'xlsx-sync-{datetime.now(timezone.utc).timestamp()}',
                    'text': comments_raw,
                    'evaluator_name': evaluator_name or 'Evaluation Committee',
                    'created_at': datetime.now(timezone.utc).isoformat()
                })

            record = {
                'category': category,
                'target_id': sub_id,
                'status': status,
                'is_next_round': is_next_round,
                'score': score_val,
                'max_score': 100,
                'last_evaluated_by_name': evaluator_name or None,
                'updated_at': datetime.now(timezone.utc).isoformat()
            }
            if comments_list:
                record['comments'] = comments_list

            records_to_upsert.append(record)

    print(f"Processed {len(records_to_upsert)} submission records across all sheets.")
    if dry_run:
        print("[DRY RUN] Sample record to upsert:")
        if records_to_upsert:
            print(records_to_upsert[0])
        return

    # Upsert to Supabase
    url = f"{SUPABASE_URL}/rest/v1/submission_evaluations?on_conflict=category,target_id"
    # Batch in chunks of 50
    batch_size = 50
    for i in range(0, len(records_to_upsert), batch_size):
        chunk = records_to_upsert[i:i + batch_size]
        res = requests.post(url, headers=HEADERS, json=chunk)
        if res.status_code in (200, 201):
            print(f"Successfully synced records {i + 1} to {min(i + batch_size, len(records_to_upsert))}")
        else:
            print(f"Error syncing batch: {res.status_code} - {res.text}")

    print("Sync complete!")

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Sync evaluation xlsx to Supabase')
    parser.add_argument('--file', default='QFF2026_Online_Competitions_Evaluation_Master.xlsx', help='Path to edited XLSX file')
    parser.add_argument('--dry-run', action='store_true', help='Preview sync payload without modifying database')
    args = parser.parse_args()

    sync_from_xlsx(file_path=args.file, dry_run=args.dry_run)
